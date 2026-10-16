"""Build the editable equipment and setup guide from the current audited basket.

Use the bundled workspace Python runtime with python-docx and Pillow installed.
This creates documentation only. It does not control or play audio hardware.
"""

from decimal import Decimal
from math import cos, pi
from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/documents/2026-10-16"
ASSETS = OUT / "assets"
AUDIT = ROOT / "sources/Sound_Chain_Qualification_Audit_2026-10-16.json"
FONT = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
SOURCES = [
    ("S1", "SONIC v1 methods", "https://doi.org/10.1101/2025.09.30.679683", "Saved PDF printed p 11 Sections 6.1 to 6.3; p 13 Section 6.6"),
    ("S2", "Earthworks M50 specifications", "https://earthworksaudio.com/measurement-microphones/m50/", "Bandwidth power load and individual calibration"),
    ("S3", "Earthworks serial calibration request", "https://earthworksaudio.com/support/ecf/", "Request the file for the delivered microphone"),
    ("S4", "Edifier MR3 specifications and delay FAQ", "https://www.edifier.com/global/p/studio-monitors/mr3", "Range inputs processing firmware and standby"),
    ("S5", "Edifier MR3 response illustration", "https://new-edifier-us-oss.edifier.com/files/20240904/9a364d5a62c02117cf79f7f52b4cc913.jpg", "Supports 32 kHz output under unspecified conditions"),
    ("S6", "Erin original MR3 measurements", "https://www.erinsaudiocorner.com/loudspeakers/edifier_mr3/", "Independent measurements end near 20 kHz"),
    ("S7", "Audient iD14 MKII specifications", "https://audient.com/products/audio-interfaces/id14/tech-specs/", "Inputs outputs power and frequency response"),
    ("S8", "Audient connections and phantom power", "https://audient.github.io/audient-documentation-site/id14/front-panel/", "Independent phantom switches and front DI behavior"),
    ("S9", "Audient output routing", "https://audient.github.io/audient-documentation-site/id14/system-panel-routing/", "DAW Thru and monitoring configuration"),
    ("S10", "REED R8090 specifications", "https://www.reedinstruments.com/pdfs/cache/www.reedinstruments.com/r8090/datasheet/r8090-datasheet.pdf", "Level accuracy frequency port and battery"),
    ("S11", "Python sounddevice playback and recording", "https://python-sounddevice.readthedocs.io/en/latest/api/convenience-functions.html", "Initial playrec checks and continuous streams"),
    ("S12", "Earthworks M50 user guide", "https://earthworksaudio.com/wp-content/uploads/2022/05/Earthworks_Audio_M50_may2022.pdf", "Included clip adapter and calibration information"),
    ("S13", "Dayton EMM-6 specifications", "https://www.parts-express.com/pedocs/manuals/390-801--dayton-audio-emm-6-quick-reference-guide.pdf", "18 Hz to 20 kHz and individual calibration file"),
    ("S14", "Edifier MR3 user manual", "https://new-edifier-us-oss.edifier.com/files/20240902/93517b14f597b1ec5b5814993aa9fca7.pdf", "Monitor mode and power rating with DRC on"),
    ("S15", "Audient iD4 MKII specifications", "https://audient.com/products/audio-interfaces/id4/tech-specs/", "Earlier one full band input comparison"),
]


def font(size):
    return ImageFont.truetype(str(FONT), size)


def diagram():
    im = Image.new("RGB", (1800, 800), "white")
    d = ImageDraw.Draw(im)
    dark, blue, gray = "#222222", "#245B87", "#D9D9D9"

    def label(x, y, text, size=31, fill=dark, anchor="mm"):
        d.multiline_text((x, y), text, font=font(size), fill=fill, anchor=anchor, align="center", spacing=10)

    def box(rect, text, size=32):
        d.rectangle(rect, outline=dark, width=3)
        label((rect[0]+rect[2])/2, (rect[1]+rect[3])/2, text, size)

    def arrow(points, fill=blue):
        d.line(points, fill=fill, width=6)
        x, y = points[-1]
        px, py = points[-2]
        dx, dy = x-px, y-py
        length = (dx*dx+dy*dy)**0.5
        ux, uy = dx/length, dy/length
        d.polygon([(x,y), (x-23*ux+10*uy,y-23*uy-10*ux), (x-23*ux-10*uy,y-23*uy+10*ux)], fill=fill)

    d.rectangle((490, 60, 970, 700), outline=gray, width=3)
    label(730, 30, "Same iD14 interface", 28)
    box((45,260,350,445), "Mac\nPython waveform\nand capture")
    box((535,100,925,275), "DAC outputs\n1 to speaker\n2 to reference")
    box((535,475,925,650), "ADC inputs\n1 microphone\n2 line reference")
    box((1240,100,1740,275), "MR3\nOne documented\ncabinet and input")
    box((1240,475,1740,650), "M50\nTip at ear\nreference position")
    arrow([(350,300),(445,300),(445,190),(535,190)])
    label(405,143,"USB\nplayback",26)
    arrow([(925,175),(1240,175)])
    label(1080,125,"Balanced TRS",28)
    arrow([(1490,275),(1490,475)])
    label(1630,375,"Sound",30)
    arrow([(1240,560),(925,560)])
    label(1080,508,"XLR cable",28)
    arrow([(700,275),(700,475)])
    label(830,375,"TRS\nreference",27)
    arrow([(535,565),(445,565),(445,405),(350,405)])
    label(390,595,"USB\nrecording",26)
    label(900,755,"R8090 and ADP1 are used for the separate level check",28)
    p = ASSETS / "sound_setup_connections.png"
    im.save(p, dpi=(300,300))
    return p


def envelope_figure():
    im = Image.new("RGB", (1800, 470), "white")
    d = ImageDraw.Draw(im)
    left, right, top, bottom = 170, 1700, 130, 350
    d.line((left,bottom,right,bottom), fill="#444444", width=3)
    points=[]
    for i in range(1001):
        ms=i/100
        value = 0.5-0.5*cos(pi*ms/2) if ms<2 else 1 if ms<=8 else 0.5-0.5*cos(pi*(10-ms)/2)
        points.append((left+(right-left)*ms/10,bottom-(bottom-top)*value))
    d.line(points,fill="#245B87",width=7)
    for ms in (0,2,8,10):
        x=left+(right-left)*ms/10
        d.line((x,bottom,x,bottom+12),fill="#444444",width=3)
        d.text((x,bottom+42),str(ms),font=font(29),fill="#222222",anchor="mm")
    for x,text in ((left+153,"2 ms rise\n192 samples"),(left+765,"6 ms plateau\n576 samples"),(right-153,"2 ms fall\n192 samples")):
        d.multiline_text((x,60),text,font=font(29),fill="#222222",anchor="mm",align="center",spacing=8)
    d.text((right,435),"Time in milliseconds",font=font(29),fill="#222222",anchor="rm")
    d.text((25,40),"Amplitude",font=font(28),fill="#222222")
    p=ASSETS / "ten_ms_envelope_schematic.png"
    im.save(p,dpi=(300,300))
    return p


def hyperlink(paragraph, text, url):
    part=paragraph.part
    rid=part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    link=OxmlElement("w:hyperlink")
    link.set(qn("r:id"),rid)
    run=OxmlElement("w:r")
    properties=OxmlElement("w:rPr")
    color=OxmlElement("w:color"); color.set(qn("w:val"),"245B87"); properties.append(color)
    under=OxmlElement("w:u"); under.set(qn("w:val"),"single"); properties.append(under)
    run.append(properties)
    value=OxmlElement("w:t"); value.text=text; run.append(value)
    link.append(run); paragraph._p.append(link)


def set_alt(run, title, description):
    for element in run._r.xpath(".//wp:docPr"):
        element.set("title",title); element.set("descr",description)


def paragraph(doc, text="", bold_lead=None, style=None):
    p=doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def heading(doc,text,level=1):
    return doc.add_heading(text,level)


def new_numbering(doc):
    numbering=doc.part.numbering_part.element
    base_id=doc.styles["List Number"].element.pPr.numPr.numId.val
    base=next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId"))==str(base_id))
    abstract_id=base.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id=max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num")))+1
    num=OxmlElement("w:num"); num.set(qn("w:numId"),str(new_id))
    abstract=OxmlElement("w:abstractNumId"); abstract.set(qn("w:val"),abstract_id); num.append(abstract)
    override=OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"),"0")
    start=OxmlElement("w:startOverride"); start.set(qn("w:val"),"1"); override.append(start); num.append(override)
    numbering.append(num)
    return new_id


def numbered_paragraph(doc,text,num_id):
    p=paragraph(doc,text,style="List Number")
    props=p._p.get_or_add_pPr().get_or_add_numPr()
    level=OxmlElement("w:ilvl"); level.set(qn("w:val"),"0"); props.append(level)
    identifier=OxmlElement("w:numId"); identifier.set(qn("w:val"),str(num_id)); props.append(identifier)
    return p


def page(doc,title):
    doc.add_page_break(); heading(doc,title)


def table(doc, headers, rows, widths, links=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for cell,width in zip(t.columns,widths): cell.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement("w:tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        el=OxmlElement("w:"+side); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"5"); el.set(qn("w:color"),"D9D9D9"); borders.append(el)
    props.append(borders)
    margins=OxmlElement("w:tblCellMar")
    for side,value in (("top",85),("bottom",85),("left",100),("right",100)):
        el=OxmlElement("w:"+side); el.set(qn("w:w"),str(value)); el.set(qn("w:type"),"dxa"); margins.append(el)
    props.append(margins)
    for j,value in enumerate(headers): t.rows[0].cells[j].text=value
    repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,value in enumerate(row):
            cells[j].text=str(value)
            if links and j==0 and i<len(links) and links[i]:
                cells[j].paragraphs[0].clear(); hyperlink(cells[j].paragraphs[0],str(value),links[i])
    for i,row in enumerate(t.rows):
        cant=OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cant)
        for j,(cell,width) in enumerate(zip(row.cells,widths)):
            cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"30343A" if i==0 else "FFFFFF" if i%2 else "F3F4F5"); cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.05
                if j>0 and len(str(row.cells[j].text))<22: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size=Pt(10.3)
                    if i==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    paragraph(doc,"").paragraph_format.space_after=Pt(0)
    return t


def picture(doc,path,width,caption,alt):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(); run.add_picture(str(path),width=Inches(width)); set_alt(run,caption,alt)
    p.paragraph_format.space_after=Pt(3)
    cp=paragraph(doc,caption); cp.style=doc.styles["Caption"]


def build():
    OUT.mkdir(parents=True,exist_ok=True); ASSETS.mkdir(exist_ok=True)
    audit=json.loads(AUDIT.read_text())
    parts=audit["revised_recommendation"]["parts"]
    subtotal=sum(Decimal(str(p["unit_usd"]))*p["quantity"] for p in parts)
    assert subtotal==Decimal("2099.87")
    doc=Document(); section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.65)
    section.footer_distance=Inches(.25)
    for name in ("Normal","Title","Subtitle","Heading 1","Heading 2","Caption"):
        s=doc.styles[name]; s.font.name="Arial"; s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),"Arial")
        s._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial")
    normal=doc.styles["Normal"]; normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.08; normal.paragraph_format.space_after=Pt(6)
    doc.styles["Title"].font.size=Pt(24)
    doc.styles["Title"].paragraph_format.space_after=Pt(5)
    doc.styles["Subtitle"].font.size=Pt(11)
    doc.styles["Heading 1"].font.size=Pt(17)
    doc.styles["Heading 1"].paragraph_format.space_before=Pt(0)
    doc.styles["Heading 1"].paragraph_format.space_after=Pt(9)
    doc.styles["Heading 2"].font.size=Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before=Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_after=Pt(4)
    doc.styles["Caption"].font.size=Pt(9)
    doc.styles["Caption"].font.bold=False
    doc.styles["Caption"].font.italic=False
    doc.styles["Caption"].paragraph_format.space_after=Pt(6)
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("M50 sound setup  |  ").font.size=Pt(9)
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)
    doc.core_properties.title="M50 auditory sound setup"
    doc.core_properties.subject="Equipment rationale wiring software calibration and qualification"
    doc.core_properties.author="Howard Wang"

    # Page 1
    doc.add_paragraph("M50 auditory sound setup",style="Title")
    doc.add_paragraph("Parts rationale wiring and verification\n16 October 2026",style="Subtitle")
    paragraph(doc,"We plan to build a Mac sound system with an Earthworks M50, Audient iD14 MKII and Edifier MR3. The first checkout will use the mouse setup, then the audio equipment will move to sheep. The M50 is the selected microphone; iD14 is the recommended interface upgrade. MR3 remains a budget speaker candidate until acoustic tests pass.")
    paragraph(doc,"Our proposed sound target is approximately 1.4 to 32 kHz, with consecutive 10 ms sine pips and 2 ms ramps at each end. We will play and record at 96 kHz. Neural recording and clock integration remain later work. [S1]")
    heading(doc,"Parts to purchase",2)
    names=["Earthworks M50","Audient iD14 MKII","Edifier MR3 powered pair","REED R8090 standard","Hosa HMIC-010 XLR 10 ft","Hosa HSS-005 TRS 5 ft","Gator RI-MICTP-FBM boom stand"]
    roles=["Measurement mic","Playback and capture","Sound delivery","1 kHz level reference","Mic to interface","Speaker and reference","Microphone placement"]
    rows=[]
    for name,role,item in zip(names,roles,parts):
        rows.append([name,item["quantity"],f"${item['unit_usd']:,.2f}",f"${item['unit_usd']*item['quantity']:,.2f}",role])
    rows.append(["Listed parts total","","","$2,099.87",""])
    table(doc,["Item and buying link","Qty","Unit","Line total","Role"],rows,[2.55,.4,.85,.95,1.45],links=[p["buy_url"] for p in parts])
    paragraph(doc,"Prices were checked on 16 October 2026. All seven listings indicated stock; none is reserved. The MR3 price is for a pair, and the TRS quantity is two cables.")
    paragraph(doc,"Budget: $2,099.87 in parts + $50 for small support or USB needs + $200 tax and shipping reserve = $2,349.87. Set aside about $2,400. The allowances are estimates, not a delivered quotation or an upper bound.",bold_lead="Budget:")
    paragraph(doc,"Included: M50 microphone clip and ADP1 calibrator adapter; iD14 USB C to USB C cable; R8090 9 V battery. MR3 has its own amplifier. We use the existing Mac. [S8, S10, S12]",bold_lead="Included:")
    paragraph(doc,"Also confirm a stable speaker shelf or support, a ruler or tape measure, position marks, a mains outlet and several GB of Mac storage. A small clip thread adapter or different USB cable may be needed. The $50 allowance assumes simple mounting or reuse; specialized mounts would add cost.")

    # Page 2
    page(doc,"Why we chose the microphone and speaker")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for fname,title in (("m50.jpg","Earthworks M50"),("edifier-mr3.jpg","Edifier MR3 pair")):
        run=p.add_run(); run.add_picture(str(ROOT/"web/assets/parts"/fname),width=Inches(1.65)); set_alt(run,title,title+" manufacturer product photograph")
        p.add_run("     ")
    paragraph(doc,"M50 on the left and MR3 pair on the right. Manufacturer product photographs; source records are retained in web/assets/parts/sources.json.",style="Caption")
    heading(doc,"Earthworks M50",2)
    paragraph(doc,"We chose the M50 because the measurement requirement extends above normal audio frequencies. Its specified 3 Hz to 50 kHz band covers our 32 kHz endpoint, and each microphone has its own response chart and electronic calibration data. This gives us a documented way to correct the microphone response when measuring the highest tones. Its XLR output connects through the Mac audio interface. [S2, S12]")
    paragraph(doc,"The higher microphone cost buys a measurement reference that we can own and reuse for mouse and sheep work. Register the delivered microphone, request its serial matched electronic calibration file, and inspect that file through at least 32 kHz. Keep the original data and its correction convention. Factory sensitivity at 1 kHz and frequency response correction are separate pieces of calibration. [S3, S12]")
    heading(doc,"Edifier MR3",2)
    paragraph(doc,"We chose the MR3 as the inexpensive speaker to evaluate. It includes amplification, accepts a balanced TRS line signal, and has a nominal 52 Hz to 40 kHz range. The manufacturer response illustration also supports output near 32 kHz, rather than relying only on a printed upper frequency. A single documented cabinet/channel keeps the initial sound field simple. [S4, S5]")
    paragraph(doc,"The remaining question is usable output at our level and geometry. The response illustration does not specify all test conditions, and published independent measurements end near 20 kHz. Neither source establishes consecutive 10 ms pips at 32 kHz. The manual rates power with dynamic range compression on; its effect on our sequence needs measurement. Keep MR3 only if the actual sequence meets the lab's level and timing criteria. [S5, S6, S14]")
    heading(doc,"Why the cheaper mic was excluded",2)
    paragraph(doc,"The Dayton EMM-6 has individual calibration data but is specified only to 20 kHz. It could support an audible subset, but cannot establish our full 32 kHz calibration requirement. The M50 choice preserves the broader sound target. [S13]")

    # Page 3
    page(doc,"Why we chose the interface and level reference")
    heading(doc,"Audient iD14 MKII",2)
    paragraph(doc,"The iD14 plays the waveform and records two signals at once: the M50 acoustic signal and a duplicate electrical waveform. Two rear microphone or line inputs support that arrangement. Its preamp at minimum gain and ADC are specified through 40 kHz; its DAC is specified through 40 kHz. It supplies 48 V phantom power at 10 mA per channel, matching the documented M50 power requirement. [S2, S7]")
    paragraph(doc,"This is why we recommend iD14 over the earlier iD4. The iD4 can perform the original acoustic checkout, but has only one full band mic or line input. The iD14 upgrade and second TRS cable add $93.96 to the previous basket and give us a stronger timing measurement path. The revised list includes one interface. Actual working gain, channel matching and native Mac playback/capture still need checks. [S7, S15]")
    paragraph(doc,"At 96 kHz, each 10 ms pip contains 960 samples. This is a functional Rice adaptation from the paper's 192 kHz DAC. The fundamental at 32 kHz is below the 48 kHz Nyquist limit, but the recording cannot capture its 64 kHz second harmonic. We can measure the required fundamentals without claiming complete ultrasonic distortion coverage. [S1, S7]")
    heading(doc,"REED R8090",2)
    paragraph(doc,"The calibrator supplies a known 94 or 114 dB signal at 1 kHz. With the M50 and a fixed interface gain, it establishes the conversion from recorded sample amplitude to sound pressure. The M50 correction file describes frequency response; the R8090 checks the recording level scale. Its stated accuracy is +/-0.5 dB, and its half inch port nominally matches the included ADP1 adapter. Verify the seal. [S10, S12]")
    paragraph(doc,"Start with the 94 dB setting and inspect headroom. The 114 dB setting can clip at a gain used for stimulus recording. The $219 standard model does not include the separately priced traceable calibration certificate. A 1 kHz reference does not independently establish accuracy at 32 kHz or certify the entire measurement chain.")
    heading(doc,"Cables and positioning",2)
    paragraph(doc,"The XLR cable carries the microphone signal and phantom power. One balanced TRS cable carries sound to the MR3; the second carries the electrical reference back to the interface. The boom stand holds the microphone tip at a repeatable reference position. It is a microphone stand, not the speaker support. Confirm the supplied clip thread fits before relying on the setup.")
    paragraph(doc,"The Mac USB connection must provide at least 0.9 A at 5 V. Prefer a direct suitable USB port with the included cable. Save the exact Mac and audio settings after confirming native 96 kHz duplex operation. [S7, S8]")
    paragraph(doc,"Software to configure: Audient iD software for interface routing, and Python with NumPy and sounddevice for waveform generation and playback/capture. Save recordings and metadata to writable Mac storage. Native operation on the chosen Mac must be checked after installation. [S8, S9, S11]")

    # Page 4
    page(doc,"How we will wire and position the equipment")
    picture(doc,diagram(),7.05,"Signal connections for simultaneous sound and electrical reference recording","Mac sends two playback channels to iD14. Output one feeds MR3. M50 measures its sound into input one. Output two returns to line input two as a duplicate electrical reference. Both inputs return to the Mac.")
    rows=[
        ["Mac to iD14","Included USB C cable","Native 96 kHz playback and capture"],
        ["Output 1 to MR3","Balanced TRS cable 1","Verify the input drives the chosen cabinet"],
        ["Output 2 to rear input 2","Balanced TRS cable 2","Line input with phantom off and low gain"],
        ["M50 to rear input 1","XLR cable","Microphone input with phantom on"],
        ["R8090 over M50 tip","Included ADP1 adapter","Separate level check then remove calibrator"],
    ]
    table(doc,["Connection","Part","Setting or check"],rows,[2.0,2.05,3.15])
    paragraph(doc,"Leave the front instrument input empty because it takes over channel 1. Mute both recorded inputs from every speaker/output mix to prevent feedback. Route the intended playback channels explicitly. DAW Thru bypasses iD Mixer level controls, so reduce source amplitude before selecting it and verify which controls affect the actual output. [S8, S9]")
    paragraph(doc,"Place the selected MR3 cabinet on a stable support and align its tweeter with the ear/reference position. Place the M50 tip at that position for the acoustic checkout, pointing as required by its free field calibration. Measure and record distance, height and orientation; photograph the arrangement and mark positions. The mouse distance is still undecided.")
    paragraph(doc,"Use wired TRS, Monitor mode and fixed documented EQ/room settings. Record firmware and standby settings; follow the manufacturer's standby guidance. The second DAC output is a duplicate reference. Characterize relative DAC/ADC channel gain and delay before treating it as a timing reference for output 1. [S4, S7 to S9]")

    # Page 5
    page(doc,"How we will program the sound")
    paragraph(doc,"Python on the Mac creates sample arrays. The interface converts them to voltage at a fixed sample rate; the speaker converts voltage to sound. We prepare each continuous block before playback so Python does not schedule individual tones using 10 ms sleeps.")
    paragraph(doc,"Select the iD14 explicitly in Python. Route macOS alerts and other applications away from it, and prevent computer sleep during a run. Save the actual device and channel mapping.")
    table(doc,["Parameter","Planned value","Meaning"],[
        ["Sample rate","96,000 samples per second","Waveform sampling resolution"],
        ["Tone duration","10 ms or 960 samples","100 consecutive pips per second"],
        ["End ramps","2 ms or 192 samples each","Included inside the 10 ms pip"],
        ["Between ramps","6 ms or 576 samples","Central plateau"],
        ["Frequency band","Approximately 1.4 to 32 kHz","Final frequency list remains a lab choice"],
    ],[1.6,2.5,3.1])
    picture(doc,envelope_figure(),7.05,"Illustrative raised cosine envelope for one 10 ms pip at 96 kHz","Schematic envelope rises over two milliseconds, remains at full amplitude for six milliseconds and falls over two milliseconds. This is a proposed ramp shape, not measured acoustic data.")
    paragraph(doc,"The paper specifies zero starting phase and 2 ms end ramps but does not specify the ramp shape. A raised cosine is our proposed implementation choice. The exact original 37 frequency array is absent from the methods; do not present a reconstructed list as that published array. [S1]")
    heading(doc,"Software sequence",2)
    sequence_id=new_numbering(doc)
    for text in [
        "Build a balanced list of the chosen frequencies, shuffle using a saved seed, and generate each sine pip using the appropriate calibrated amplitude.",
        "Join the pips without silence between them. For timing tests, retain leading and trailing silence so acoustic onset and final decay are captured.",
        "Duplicate the waveform to playback outputs 1 and 2. Record microphone input 1 and electrical reference input 2 at the same native sample rate.",
        "Use sounddevice playrec for short initial checks and an explicit continuous duplex stream for controlled longer runs. Inspect underflow/overflow status. [S11]",
    ]:
        numbered_paragraph(doc,text,sequence_id)
    paragraph(doc,"Save the frequency schedule, digital onset sample numbers, random seed, amplitude table, exact waveform or its reproducible source, raw two channel recording and all device/geometry settings. Keep raw data alongside any corrected result. A 30 minute stereo float32 capture at 96 kHz is about 1.38 GB; stream blocks to disk rather than requiring the whole session in memory.")

    # Page 6
    page(doc,"Calibration and the tests that qualify the setup")
    heading(doc,"Establish the measurement scale",2)
    calibration_id=new_numbering(doc)
    for text in [
        "Register the M50 and request its serial matched ECF. Inspect coverage through 32 kHz, units and whether values describe microphone response or corrective gain. Confirm the correct interpretation before applying it. [S2, S3]",
        "Set a microphone gain, fit the ADP1 and R8090, and record the 94 dB reference. Check the seal and avoid clipping. Remove the calibrator for free field measurement. Keep the gain fixed; recalibrate after any gain change. [S10, S12]",
        "At the empty ear/reference position, record background sound and test tones across the selected frequencies. Apply the microphone correction and fixed gain level scale to estimate sound pressure. Check the interface response at the operating settings as well.",
        "Build a per frequency playback amplitude table that reaches the selected level within measured speaker headroom. Define whether SPL refers to the central plateau or RMS across the whole ramped pip; use the same convention for calibration and validation.",
    ]:
        numbered_paragraph(doc,text,calibration_id)
    heading(doc,"Measure the actual sequence",2)
    table(doc,["Check","Measurement","Pass definition"],[
        ["Frequency and level","All retained tones including 27 and 32 kHz","Lab level and tolerance at fixed geometry"],
        ["Short pip behavior","Actual 10 ms pips and large frequency jumps","Duration spacing and residual energy tolerances"],
        ["Timing reference","Microphone versus characterized electrical reference","Delay and variability within lab criteria"],
        ["Stability","Repeated complete blocks and before/after level checks","No unexplained missing events clipping or drift"],
    ],[1.45,3.0,2.75])
    paragraph(doc,"Paper blocks contain 8,000 to 15,000 pips, lasting 80 to 150 seconds. At least 12 blocks imply about 16 to 30 minutes plus 1 second block gaps. This is a useful paper style duration test, not an automatically approved mouse protocol. [S1]")
    paragraph(doc,"A fixed playback delay longer than 10 ms can still preserve a 100 pip per second sequence. Frequency dependent delay, compression and ringing need acoustic recordings. Sample arithmetic and long steady tones alone do not establish this behavior.")
    paragraph(doc,"Keep the MR3 only after it meets the agreed acoustic criteria. A quick mouse checkout is the first stage. When moving the equipment to sheep, repeat level/frequency calibration and sequence checks at the new geometry. Neural marker/clock alignment must be designed separately before neural timing claims.")

    # Page 7
    page(doc,"Decisions and source references")
    paragraph(doc,"The M50 is selected. iD14 and the extra reference cable are the recommended verification revision. The remaining decisions belong to the experimental team, and the current list has not passed a physical acoustic test.")
    table(doc,["Decision","Current status","Owner"],[
        ["Mouse distance and orientation","TBD before acoustic checkout","Howard and lab"],
        ["Mouse level and SPL convention","TBD under the lab protocol","Lab lead with Howard"],
        ["Frequency list and ramp shape","TBD; raised cosine proposed","Howard and lab"],
        ["Numerical level and timing tolerances","TBD before judging measurements","Lab lead with Howard"],
        ["Sheep geometry and neural alignment","Later stage; recalibration required","Howard Jiaao and lab"],
    ],[2.3,3.2,1.7])
    paragraph(doc,"SONIC reports awake, alert sheep and approximately 80 dB SPL at sheep ear distance. Planned Rice sheep sedation is a team adaptation. The paper level is not a mouse exposure prescription, and neither the hardware list nor an acoustic checkout establishes neural performance. [S1]")
    heading(doc,"Sources",2)
    for key,title,url,scope in SOURCES:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
        p.add_run(key+"  ").bold=True
        hyperlink(p,title,url)
        p.add_run(". "+scope+".")
        for run in p.runs: run.font.size=Pt(9.5)
    paragraph(doc,"Buying links are embedded in the parts table on page 1. Manufacturer product photographs are reused from the project's recorded image sources. Wiring and envelope figures are explanatory schematics, not experimental measurements.",style="Caption")
    dest=OUT/"M50_Sound_Setup_Guide.docx"
    doc.save(dest)
    manifest={
        "date":"2026-10-16","document":str(dest.relative_to(ROOT)),
        "source_audit":str(AUDIT.relative_to(ROOT)),"intended_pages":7,
        "parts_total_usd":str(subtotal),"working_budget_usd":"2349.87",
        "photo_provenance":"web/assets/parts/sources.json",
        "source_references":[dict(id=k,title=t,url=u,scope=s) for k,t,u,s in SOURCES],
        "scientific_scope":"Documentation and setup plan only; physical qualification pending; initial mouse checkout then independently calibrated sheep geometry",
        "software_scope":"Architecture and workflow described; no stimulus control implementation or audio playback performed",
    }
    (ROOT/"sources/M50_Sound_Setup_Guide_2026-10-16.json").write_text(json.dumps(manifest,indent=2)+"\n")
    print(dest)


if __name__=="__main__":
    build()
